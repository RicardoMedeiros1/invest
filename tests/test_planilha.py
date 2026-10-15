"""Testes da planilha única (carteira.xlsx): sempre o mesmo arquivo, abas substituídas sem apagar as outras."""
import numpy as np
import openpyxl
import pandas as pd
import pytest

import planilha as pl


def df(**colunas):
    return pd.DataFrame(colunas)


def ler(caminho):
    return pd.read_excel(caminho, sheet_name=None)


def test_cria_a_planilha_com_uma_aba_por_tabela(tmp_path):
    caminho = tmp_path / "carteira.xlsx"
    pl.atualizar_planilha(caminho, {"Ranking": df(ativo=["PETR4", "VALE3"], score=[0.957, 0.879]),
                                    "Comparação": df(nome=["Estratégia"], xirr=[20.68])})
    abas = ler(caminho)
    assert list(abas) == ["Comparação", "Ranking"] or set(abas) == {"Ranking", "Comparação"}
    assert abas["Ranking"].ativo.tolist() == ["PETR4", "VALE3"] and abas["Comparação"].xirr.tolist() == [20.68]


def test_atualizar_substitui_so_as_abas_dadas_e_mantem_as_outras(tmp_path):
    caminho = tmp_path / "carteira.xlsx"
    pl.atualizar_planilha(caminho, {"Melhores agora": df(a=[1]), "Ranking": df(a=[1, 2, 3])})
    pl.atualizar_planilha(caminho, {"Ranking": df(a=[9]), "Backtest": df(b=[5])})            # outro comando roda depois
    abas = ler(caminho)
    assert set(abas) == {"Melhores agora", "Ranking", "Backtest"}                          # nada duplicado nem perdido
    assert abas["Melhores agora"].a.tolist() == [1]                                         # a aba que não foi tocada continua
    assert abas["Ranking"].a.tolist() == [9]                                                # a tocada foi substituída por inteiro


def test_so_existe_um_arquivo_mesmo_depois_de_varias_execucoes(tmp_path):
    caminho = tmp_path / "carteira.xlsx"
    for i in range(4):
        pl.atualizar_planilha(caminho, {"Ranking": df(a=[i])}, historico={"rodada": i})
    assert sorted(p.name for p in tmp_path.iterdir()) == ["carteira.xlsx"]


def test_abas_ficam_na_ordem_combinada(tmp_path):
    caminho = tmp_path / "carteira.xlsx"
    pl.atualizar_planilha(caminho, {"Operações": df(a=[1]), "Ranking": df(a=[1]), "Outra": df(a=[1]), "Melhores agora": df(a=[1])},
                          historico={"x": 1})
    assert openpyxl.load_workbook(caminho).sheetnames == ["Melhores agora", "Ranking", "Operações", "Histórico", "Outra"]
    pl.atualizar_planilha(caminho, {"Comparação": df(a=[1])})                                # nova aba entra no lugar certo
    assert openpyxl.load_workbook(caminho).sheetnames == ["Melhores agora", "Ranking", "Comparação", "Operações", "Histórico", "Outra"]


def test_historico_acumula_uma_linha_por_execucao(tmp_path):
    caminho = tmp_path / "carteira.xlsx"
    pl.atualizar_planilha(caminho, {"Ranking": df(a=[1])}, historico={"configuracao": "só ações", "retorno": 20.68})
    pl.atualizar_planilha(caminho, {"Ranking": df(a=[2])}, historico={"configuracao": "3 + 2", "retorno": 18.08})
    pl.atualizar_planilha(caminho, {"Melhores agora": df(a=[1])})                           # o painel simples não mexe no histórico
    pl.atualizar_planilha(caminho, {"Ranking": df(a=[3])}, historico={"configuracao": "4 + 1", "retorno": 17.43})
    h = ler(caminho)["Histórico"]
    assert h.configuracao.tolist() == ["só ações", "3 + 2", "4 + 1"] and h.retorno.tolist() == [20.68, 18.08, 17.43]


def test_historico_entra_numa_planilha_que_so_tinha_o_painel(tmp_path):
    caminho = tmp_path / "carteira.xlsx"
    pl.atualizar_planilha(caminho, {"Melhores agora": df(a=[1])})                           # sem aba Histórico ainda
    pl.atualizar_planilha(caminho, {"Ranking": df(a=[1])}, historico={"retorno": 1.5})
    assert ler(caminho)["Histórico"].retorno.tolist() == [1.5]


def test_numeros_saem_arredondados_sem_mexer_no_original(tmp_path):
    caminho = tmp_path / "carteira.xlsx"
    original = df(preco=[49.123456789], score=[0.9571234], n=[7], nome=["x"])
    pl.atualizar_planilha(caminho, {"Ranking": original})
    r = ler(caminho)["Ranking"]
    assert r.preco.iloc[0] == 49.12 and r.score.iloc[0] == 0.957 and r.n.iloc[0] == 7
    assert original.preco.iloc[0] == 49.123456789                                           # o DataFrame do chamador não muda


def test_cabecalho_em_negrito_linha_fixa_e_colunas_ajustadas(tmp_path):
    caminho = tmp_path / "carteira.xlsx"
    pl.atualizar_planilha(caminho, {"Corretoras": df(nome=["Nubank"], atencao=["x" * 300])})
    ws = openpyxl.load_workbook(caminho)["Corretoras"]
    assert ws["A1"].font.bold and ws["B1"].font.bold and not ws["A2"].font.bold
    assert ws.freeze_panes == "A2"
    assert ws.column_dimensions["A"].width < ws.column_dimensions["B"].width <= 60


def test_aba_vazia_nao_quebra(tmp_path):
    caminho = tmp_path / "carteira.xlsx"
    pl.atualizar_planilha(caminho, {"Operações": pd.DataFrame(), "Ranking": df(a=[1])})
    assert "Operações" in openpyxl.load_workbook(caminho).sheetnames


def test_nome_da_planilha_aceita_com_ou_sem_extensao():
    assert str(pl.caminho_da_planilha("carteira")) == "carteira.xlsx"
    assert str(pl.caminho_da_planilha("carteira.xlsx")) == "carteira.xlsx"
    assert str(pl.caminho_da_planilha("pasta/carteira.XLSX")).lower() == "pasta/carteira.xlsx"
    assert str(pl.caminho_da_planilha("minha.carteira")) == "minha.carteira.xlsx"


def test_planilha_aberta_no_excel_gera_mensagem_clara(tmp_path, monkeypatch):
    caminho = tmp_path / "carteira.xlsx"
    pl.atualizar_planilha(caminho, {"Ranking": df(a=[1])})

    def bloqueado(*a, **k):
        raise PermissionError(13, "Permission denied")
    monkeypatch.setattr(pl.pd, "ExcelWriter", bloqueado)
    with pytest.raises(SystemExit) as erro:
        pl.atualizar_planilha(caminho, {"Ranking": df(a=[2])})
    assert "carteira.xlsx" in str(erro.value) and "Excel" in str(erro.value) and "feche" in str(erro.value)


def test_arquivo_danificado_gera_mensagem_clara(tmp_path):
    caminho = tmp_path / "carteira.xlsx"
    caminho.write_bytes(b"isto nao e uma planilha")
    with pytest.raises(SystemExit, match="danificado"):
        pl.atualizar_planilha(caminho, {"Ranking": df(a=[1])})
    with pytest.raises(SystemExit, match="danificado"):
        pl.atualizar_planilha(caminho, {"Ranking": df(a=[1])}, historico={"x": 1})
    assert caminho.read_bytes() == b"isto nao e uma planilha"                                # não apaga nada sem avisar
