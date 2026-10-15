# -*- coding: utf-8 -*-
"""Uma única planilha (carteira.xlsx), sempre o mesmo arquivo, com uma aba para cada coisa.

Cada execução substitui só as abas que ela produz e preserva as demais. A aba Histórico acumula uma linha
por execução do scanner, para comparar versões (por exemplo, com e sem FIIs) sem criar novos arquivos.
"""
from __future__ import annotations

import zipfile
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

ABA_HISTORICO = "Histórico"
# Ordem das abas: primeiro as do painel simples, depois as do scanner completo.
ORDEM_ABAS = ["Melhores agora", "Quanto subiu", "Plano de compra", "Passado", "Corretoras", "Todos os ativos",
              "Ranking", "Diagnóstico", "Comparação", "Backtest", "Operações", ABA_HISTORICO]


def caminho_da_planilha(saida: str) -> Path:
    """'carteira' -> carteira.xlsx (e 'carteira.xlsx' continua carteira.xlsx)."""
    p = Path(saida)
    return p if p.suffix.lower() == ".xlsx" else p.with_name(p.name + ".xlsx")


def _arredondar(df: pd.DataFrame) -> pd.DataFrame:
    """Números com no máximo 2 casas (3 nas notas), para a planilha não mostrar 15 dígitos."""
    df = df.copy()
    for col in df.columns:
        if pd.api.types.is_float_dtype(df[col]):
            df[col] = df[col].round(3 if str(col).lower() in ("score", "nota") else 2)
    return df


def _formatar(ws, df: pd.DataFrame) -> None:
    """Cabeçalho em negrito, primeira linha fixa e largura das colunas ajustada ao conteúdo."""
    if df.shape[1] == 0:
        return
    for celula in ws[1]:
        celula.font = Font(bold=True)
    ws.freeze_panes = "A2"
    for i, col in enumerate(df.columns, start=1):
        comprimento = max([len(str(col))] + [len(str(v)) for v in df[col].head(200)])
        ws.column_dimensions[get_column_letter(i)].width = min(max(comprimento + 2, 8), 60)


def _ordenar(book) -> None:
    posicao = {nome: i for i, nome in enumerate(ORDEM_ABAS)}
    book._sheets.sort(key=lambda ws: posicao.get(ws.title, len(posicao)))   # estável: abas desconhecidas ficam no fim
    book.active = 0


def atualizar_planilha(caminho: Path, abas: dict[str, pd.DataFrame], historico: dict | None = None) -> Path:
    """Grava `abas` em `caminho`, substituindo só essas abas e mantendo as outras.

    Com `historico` (um dicionário), acrescenta uma linha à aba Histórico em vez de substituí-la.
    Se a planilha estiver aberta no Excel, o Windows impede a gravação: nesse caso o programa para com
    uma mensagem em português.
    """
    caminho = Path(caminho)
    abas = dict(abas)
    modo = "a" if caminho.exists() else "w"
    try:
        if historico is not None:
            antigo = None
            if modo == "a":
                try:
                    antigo = pd.read_excel(caminho, sheet_name=ABA_HISTORICO)
                except ValueError:       # a planilha existe, mas ainda não tem a aba Histórico
                    antigo = None
            nova = pd.DataFrame([historico])
            abas[ABA_HISTORICO] = nova if antigo is None else pd.concat([antigo, nova], ignore_index=True)
        opcoes = {"if_sheet_exists": "replace"} if modo == "a" else {}
        with pd.ExcelWriter(caminho, engine="openpyxl", mode=modo, **opcoes) as escritor:
            for nome, df in abas.items():
                df = _arredondar(df)
                df.to_excel(escritor, sheet_name=nome, index=False)
                _formatar(escritor.sheets[nome], df)
            _ordenar(escritor.book)
    except PermissionError as exc:
        raise SystemExit(f"Não consegui gravar {caminho}. Se ele estiver aberto no Excel, feche-o e rode de novo.") from exc
    except zipfile.BadZipFile as exc:
        raise SystemExit(f"O arquivo {caminho} parece danificado. Renomeie ou apague e rode de novo.") from exc
    return caminho
